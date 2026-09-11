# -*- coding: utf-8 -*-
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DOCX = Path(r"C:\Users\LENOVO\Desktop\软著\v3高速车辆轨迹规划跟车软件使用说明书(6).docx")
OUT_DIR = ROOT / "outputs" / "soft_copyright"
ASSET_DIR = ROOT / "outputs" / "soft_copyright_assets"
ORIG_IMG_DIR = ASSET_DIR / "original_images"
OUT_DOCX = OUT_DIR / "高速车辆轨迹规划跟车软件V1.0_软著说明书_修订版.docx"

SOFTWARE_NAME = "高速车辆轨迹规划跟车软件 V1.0"
HEADER_TEXT = "高速车辆轨迹规划跟车软件 使用说明书 V1.0"


def ensure_dirs():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    ORIG_IMG_DIR.mkdir(parents=True, exist_ok=True)


def extract_original_images():
    if not SOURCE_DOCX.exists():
        return []
    with ZipFile(SOURCE_DOCX) as zf:
        media = [name for name in zf.namelist() if name.startswith("word/media/")]
        for name in media:
            target = ORIG_IMG_DIR / Path(name).name
            target.write_bytes(zf.read(name))
    return sorted(ORIG_IMG_DIR.glob("image*.png"))


def font_path():
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for item in candidates:
        if item.exists():
            return str(item)
    return None


def make_font(size, bold=False):
    fp = font_path()
    if fp:
        return ImageFont.truetype(fp, size=size)
    return ImageFont.load_default()


def draw_wrapped(draw, text, box, font, fill=(20, 32, 46), line_gap=8, align="center"):
    x1, y1, x2, y2 = box
    width = x2 - x1
    lines = []
    for raw_line in str(text).splitlines():
        current = ""
        for ch in raw_line:
            candidate = current + ch
            bbox = draw.textbbox((0, 0), candidate, font=font)
            if bbox[2] - bbox[0] <= width or not current:
                current = candidate
            else:
                lines.append(current)
                current = ch
        if current:
            lines.append(current)
    if not lines:
        lines = [""]
    sample = draw.textbbox((0, 0), "高速车辆轨迹规划", font=font)
    line_h = sample[3] - sample[1] + line_gap
    total_h = line_h * len(lines) - line_gap
    y = y1 + max(0, (y2 - y1 - total_h) // 2)
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        if align == "center":
            x = x1 + (width - (bbox[2] - bbox[0])) // 2
        else:
            x = x1
        draw.text((x, y), line, font=font, fill=fill)
        y += line_h


def rounded_box(draw, xy, fill, outline, text, font, text_fill=(20, 32, 46)):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=2)
    pad = 18
    draw_wrapped(draw, text, (xy[0] + pad, xy[1] + pad, xy[2] - pad, xy[3] - pad), font, text_fill)


def arrow(draw, start, end, fill=(50, 92, 130), width=4):
    draw.line([start, end], fill=fill, width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        sign = 1 if ex >= sx else -1
        points = [(ex, ey), (ex - sign * 14, ey - 8), (ex - sign * 14, ey + 8)]
    else:
        sign = 1 if ey >= sy else -1
        points = [(ex, ey), (ex - 8, ey - sign * 14), (ex + 8, ey - sign * 14)]
    draw.polygon(points, fill=fill)


def make_diagrams():
    title_font = make_font(32)
    box_font = make_font(17)
    small_font = make_font(16)
    bg = (247, 250, 252)
    blue = (222, 238, 252)
    green = (224, 245, 235)
    gold = (255, 243, 213)
    purple = (236, 229, 250)
    border = (83, 112, 146)

    # Architecture diagram
    img = Image.new("RGB", (1600, 980), bg)
    d = ImageDraw.Draw(img)
    d.text((60, 42), "软件总体架构示意图", font=title_font, fill=(13, 45, 77))
    boxes = {
        "entry": (80, 130, 410, 270, blue, "主入口 run.py\n命令行参数解析\nYAML 配置加载"),
        "ui": (80, 340, 410, 480, purple, "可视化界面\npygame_dashboard.py\nPygame/Tk 后端"),
        "cfg": (520, 130, 850, 270, gold, "配置文件\n目录 tools/cfgs\n算法与环境参数"),
        "rl": (520, 340, 850, 480, green, "强化学习算法\nDDPG/TRPO/A2C/PPO2/SAC\nStable Baselines"),
        "env": (960, 130, 1310, 270, blue, "CARLA Gym 环境\nCarlaGymEnv-v1/v2\n仿真状态与奖励"),
        "planner": (960, 340, 1310, 480, green, "Frenet 轨迹规划\n候选轨迹采样\n选定轨迹输出"),
        "control": (960, 535, 1310, 675, gold, "车辆控制\nPID 控制器\nIDM 跟驰模型"),
        "carla": (960, 730, 1310, 870, purple, "CARLA 服务端\nTown04 场景\n车辆/交通流/传感器"),
        "out": (520, 730, 850, 870, blue, "输出持久化\nlogs/agent_x\nmodels / trajectories\nmonitor.csv"),
        "plot": (80, 730, 410, 870, green, "绘图与导出\n奖励曲线/轨迹/对比图\noutputs/pygame_ui\noutputs/exports"),
    }
    for _, (x1, y1, x2, y2, fill, text) in boxes.items():
        rounded_box(d, (x1, y1, x2, y2), fill, border, text, box_font)
    for s, e in [
        ((410, 200), (520, 200)),
        ((850, 200), (960, 200)),
        ((685, 270), (685, 340)),
        ((850, 410), (960, 410)),
        ((1135, 270), (1135, 340)),
        ((1135, 480), (1135, 535)),
        ((1135, 675), (1135, 730)),
        ((960, 800), (850, 800)),
        ((520, 800), (410, 800)),
        ((245, 480), (245, 730)),
        ((245, 340), (245, 270)),
    ]:
        arrow(d, s, e)
    d.text((90, 920), "说明：训练/测试由 run.py 发起；Dashboard 读取 logs 目录中的监控、模型和轨迹文件并生成图像或导出归档。", font=small_font, fill=(54, 67, 82))
    arch_path = ASSET_DIR / "software_architecture.png"
    img.save(arch_path)

    # Runtime flow diagram
    img = Image.new("RGB", (1600, 880), bg)
    d = ImageDraw.Draw(img)
    d.text((60, 42), "软件运行流程示意图", font=title_font, fill=(13, 45, 77))
    flow = [
        ("准备 CARLA 服务端\n端口 2000\nTM 端口 8000", blue),
        ("编辑 YAML 配置\n选择算法\n交通密度与规划参数", gold),
        ("执行 run.py\n--cfg_file\n--agent_id\n--env", green),
        ("创建 Gym 环境\n连接 CARLA\n加载 Town04", purple),
        ("训练或测试策略\n记录奖励\n轨迹和模型", blue),
        ("绘图/查看/导出\nDashboard\nmonitor_plot.py", green),
    ]
    y = 190
    prev_center = None
    for i, (text, fill) in enumerate(flow):
        x1 = 75 + i * 250
        x2 = x1 + 220
        rounded_box(d, (x1, y, x2, y + 175), fill, border, text, box_font)
        d.ellipse((x1 + 83, y - 68, x1 + 137, y - 14), fill=(13, 82, 128))
        d.text((x1 + 99, y - 61), str(i + 1), font=make_font(24), fill=(255, 255, 255))
        center = (x1, y + 77)
        if prev_center:
            arrow(d, (prev_center[0] + 220, prev_center[1] + 10), (x1 - 10, center[1] + 10))
        prev_center = (x1, y + 77)
    note_boxes = [
        (180, 505, 520, 640, "训练输出\nmonitor.csv\nreproduction_info.txt\n配置副本、models/*.zip"),
        (625, 505, 965, 640, "轨迹输出\ntrajectory_*.csv\nplans/latest_plans.csv\nsurrounding/*.csv"),
        (1070, 505, 1410, 640, "图像输出\nrewards_*.png\ntrajectory_*.png\ncomparison_*.png\nartifact_overview_*.png"),
    ]
    for xy in note_boxes:
        rounded_box(d, xy[:4], (255, 255, 255), (190, 204, 220), xy[4], small_font)
    flow_path = ASSET_DIR / "software_runtime_flow.png"
    img.save(flow_path)
    return arch_path, flow_path


def set_cell_text(cell, text, bold=False, color=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_table_borders(table, color="B8C2CC"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.first_child_found_in("w:tblCellMar")
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tbl_cell_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_cell_margins(table)
    for idx, h in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, "F2F4F7")
        set_cell_text(cell, h, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], str(value))
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def add_para(doc, text="", style=None, bold=False):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(80, 92, 108)
    return p


def add_picture(doc, path, caption, width=6.25):
    if path and Path(path).exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        add_caption(doc, caption)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def setup_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.10

    for list_style in ["List Bullet", "List Number"]:
        st = styles[list_style]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = HEADER_TEXT
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].font.name = "Microsoft YaHei"
    header.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(90, 100, 110)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    add_page_number(footer)
    run = footer.add_run(" 页")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_cover(doc):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("高速车辆轨迹规划跟车软件")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(26)
    run.bold = True
    run.font.color.rgb = RGBColor(13, 45, 77)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("使用说明书 / 软件设计说明书")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(56, 70, 86)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("V1.0")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(16)
    run.bold = True
    for _ in range(10):
        doc.add_paragraph()
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["软件名称", SOFTWARE_NAME],
            ["文档类型", "计算机软件著作权登记用软件说明书"],
            ["修订日期", "2026-05-26"],
            ["原始版本作者", "彭靖"],
        ],
        widths=[1.6, 4.8],
    )
    doc.add_page_break()


def add_overview(doc, flow_path):
    doc.add_heading("1 软件概述", level=1)
    doc.add_heading("1.1 软件名称与版本", level=2)
    add_para(doc, "软件全称统一为“高速车辆轨迹规划跟车软件 V1.0”。本文档中的封面、页眉、正文、图题、表题和附录均采用该名称或其不含版本号的简称“高速车辆轨迹规划跟车软件”。")
    doc.add_heading("1.2 软件用途", level=2)
    add_para(doc, "本软件面向 CARLA 高速道路仿真场景，用于在 OpenAI Gym 接口下执行强化学习策略训练、测试和 Frenet 局部轨迹规划。用户可通过命令行配置训练或测试任务，也可通过可视化 Dashboard 查看奖励曲线、轨迹文件、周围交通轨迹、方法对比图和训练产物索引。")
    doc.add_heading("1.3 适用对象", level=2)
    for item in [
        "需要在 CARLA 仿真环境中运行车辆跟车、变道、避障与轨迹跟踪实验的开发或测试人员。",
        "需要查看强化学习训练日志、模型权重、轨迹 CSV 和奖励曲线的算法调试人员。",
        "需要根据项目源码审查软件功能边界、输入输出和运行流程的软件著作权登记审查人员。",
    ]:
        add_bullet(doc, item)
    doc.add_heading("1.4 软件功能概述", level=2)
    add_table(
        doc,
        ["功能类别", "代码依据", "功能说明"],
        [
            ["仿真连接", "run.py、tools/modules.py、carla_gym/envs", "通过 carla_host、carla_port、tm_port 参数连接 CARLA 服务端，加载 Town04 场景并创建交通流。"],
            ["算法训练/测试", "run.py、tools/cfgs/*.yaml", "支持 DDPG、TRPO、A2C、PPO2、SAC 算法的训练与模型加载测试；SAC 可根据配置启用课程训练。"],
            ["Frenet 规划", "agents/local_planner", "根据全局路线和车辆状态生成候选轨迹并选取执行轨迹。"],
            ["车辆控制", "agents/low_level_controller", "使用 PID 控制器和 IDM 跟驰模型执行车辆控制。"],
            ["绘图与导出", "monitor_plot.py、tools/pygame_dashboard.py", "读取 monitor.csv、trajectory CSV 和模型目录，生成奖励曲线、轨迹图、交通轨迹图、对比图和产物概览图。"],
            ["结果保存", "logs/、outputs/", "保存训练日志、配置副本、模型权重、轨迹 CSV、绘图 PNG 和导出归档。"],
        ],
        widths=[1.25, 1.75, 3.4],
    )
    doc.add_heading("1.5 软件运行流程概述", level=2)
    add_para(doc, "软件运行时先启动 CARLA 服务端，再通过 YAML 文件和命令行参数确定算法、环境、交通密度、渲染方式和训练步数。run.py 创建 Gym 环境并连接 CARLA，训练或测试过程产生日志、模型和轨迹文件。Dashboard 与绘图脚本再读取这些文件生成图像或导出归档。")
    add_picture(doc, flow_path, "图1-1 软件运行流程示意图。该图说明 CARLA 服务端、YAML 配置、run.py、Gym 环境、训练/测试和结果绘图导出之间的执行关系。")


def add_environment(doc):
    doc.add_heading("2 运行环境", level=1)
    doc.add_heading("2.1 硬件环境", level=2)
    add_para(doc, "项目 README 提示 CARLA 对 GPU 性能要求较高。建议使用 64 位处理器、16 GB 及以上内存，并配置支持 CARLA 运行的独立显卡。更高分辨率、更多交通车辆或 3D 渲染模式会增加显卡和 CPU 负载。")
    doc.add_heading("2.2 操作系统环境", level=2)
    add_para(doc, "README 中给出的 CARLA 服务端示例主要为 Linux 命令；原说明书同时给出 Windows 运行环境。结合当前工作区实际路径，本软件可在 Windows 开发环境中编辑和运行 Python 代码，CARLA 服务端命令应以本机安装的 CARLA 版本为准。")
    doc.add_heading("2.3 Python 与依赖库环境", level=2)
    add_para(doc, "requirements.txt 中列出的主要依赖包括 tensorflow==1.14、future、numpy、pygame、pandas、mpi4py、gitpython、gym、scipy、joblib、cloudpickle、opencv-python、matplotlib、easydict、pyyaml。README 要求 Python 3.7 或更新版本，当前项目中存在 .venv 虚拟环境目录。")
    doc.add_heading("2.4 CARLA 仿真环境", level=2)
    add_para(doc, "README 推荐使用 CARLA 0.9.9.2。run.py 默认连接 127.0.0.1:2000，Traffic Manager 默认端口为 8000；tools/modules.py 中 ModuleWorld 通过 carla.Client 建立连接，加载 Town04 场景并设置 ClearNoon 天气。")
    doc.add_heading("2.5 目录与文件准备", level=2)
    add_table(
        doc,
        ["项目路径", "类型", "用途"],
        [
            ["run.py", "入口脚本", "训练、测试、模型加载、环境初始化和算法选择。"],
            ["tools/cfgs/", "配置目录", "存放 DDPG、TRPO、A2C、PPO2、SAC 及 SAC 变体配置文件。"],
            ["carla_gym/envs/", "环境模块", "定义 CarlaGymEnv-v1/v2，与 CARLA 交互并记录轨迹。"],
            ["agents/local_planner/", "规划模块", "Frenet 候选轨迹生成与轨迹转换。"],
            ["agents/low_level_controller/", "控制模块", "PID 控制器与 IDM 跟驰模型。"],
            ["logs/agent_x/", "日志目录", "保存 monitor.csv、配置副本、reproduction_info.txt、models 和 trajectories。"],
            ["outputs/pygame_ui/", "图像目录", "保存 Dashboard 生成的 PNG 预览和绘图。"],
            ["outputs/exports/", "导出目录", "保存 Export 按钮导出的日志、模型、轨迹和图像归档。"],
        ],
        widths=[1.75, 1.0, 3.65],
    )


def add_structure(doc, arch_path):
    doc.add_heading("3 软件结构与功能模块", level=1)
    doc.add_heading("3.1 软件总体架构", level=2)
    add_para(doc, "软件由命令行入口、配置文件、CARLA Gym 环境、强化学习算法库、Frenet 轨迹规划、车辆控制、日志/模型保存和可视化 Dashboard 组成。各模块通过文件路径、命令行参数和 Gym 环境接口进行调用。")
    add_picture(doc, arch_path, "图3-1 软件总体架构示意图。该图展示 run.py、配置文件、强化学习算法、CARLA Gym 环境、Frenet 规划、车辆控制和输出模块之间的关系。")
    rows = [
        ["3.2 仿真环境连接模块", "读取 --carla_host、--carla_port、--tm_port 和 --carla_res。", "创建 carla.Client，加载 Town04，初始化世界、交通流、HUD 和输入模块。", "CARLA 世界对象、车辆 Actor、交通流对象。"],
        ["3.3 参考路径与场景配置模块", "road_maps/global_route_town04.npy、CARLA 地图和 YAML 环境字段。", "读取或生成全局路线，将路线传给 Frenet Planner 与交通模块。", "全局参考路线、地图图像与路点数据。"],
        ["3.4 强化学习算法配置模块", "tools/cfgs/*.yaml、命令行 --cfg_file。", "根据 POLICY.NAME 实例化 DDPG、TRPO、A2C、PPO2 或 SAC；根据配置设置网络类型和 SAC 参数。", "训练模型对象和算法配置副本。"],
        ["3.5 Frenet 轨迹规划模块", "车辆 Frenet 状态、目标速度、车道变化动作。", "计算候选 Frenet 路径、转换为世界坐标并选择执行轨迹。", "候选轨迹和 selected 轨迹 CSV。"],
        ["3.6 车辆控制模块", "规划路径、车辆姿态、速度误差和前车信息。", "使用 VehiclePIDController 和 IntelligentDriverModel 产生控制量。", "车辆控制命令、仿真车辆轨迹。"],
        ["3.7 训练执行与监控模块", "算法配置、环境状态、训练步数。", "执行 learn 或 predict 循环，Monitor 记录 episode 奖励、长度和时间。", "monitor.csv、终端日志、模型文件。"],
        ["3.8 轨迹绘制与奖励曲线绘制模块", "monitor.csv、trajectory_*.csv、plans/*.csv。", "Dashboard 或 monitor_plot.py 绘制奖励曲线、轨迹图和交通轨迹图。", "rewards_*.png、trajectory_*.png、surrounding_tracks_*.png。"],
        ["3.9 多算法对比分析模块", "多个 agent_id 的日志目录。", "统计尾部平均奖励、最大奖励、模型数量、轨迹文件数量并绘图。", "comparison_*.png、artifact_overview_*.png。"],
        ["3.10 数据、图像、日志和模型保存模块", "训练过程数据、配置文件、图像输出。", "创建 logs/agent_x、models、trajectories、outputs/pygame_ui 和 outputs/exports。", "zip 模型、CSV、PNG、导出归档。"],
        ["3.11 参数重置与退出模块", "程序退出事件、ESC 键、窗口关闭事件。", "Dashboard 刷新数据源；run.py finally 块保存模型并调用 env.destroy。", "环境销毁、世界设置恢复、程序退出。"],
    ]
    add_table(doc, ["模块", "输入", "处理过程", "输出"], rows, widths=[1.45, 1.55, 2.0, 1.4])


def add_installation(doc, img_paths):
    doc.add_heading("4 软件安装与启动", level=1)
    doc.add_heading("4.1 启动 CARLA 服务端", level=2)
    add_para(doc, "在 CARLA 安装目录启动服务端。README 示例命令如下，实际可执行文件名和路径以本机 CARLA 安装包为准：")
    add_para(doc, "./CarlaUE4.sh -carla-server -fps=20 -world-port=2000 -windowed -ResX=1280 -ResY=720 -carla-no-hud -quality-level=Low")
    add_picture(doc, img_paths.get("carla"), "图4-1 CARLA 仿真服务端界面。该图显示 CARLA 场景窗口，软件通过默认 2000 端口与该服务端建立连接。")
    doc.add_heading("4.2 激活 Python 虚拟环境", level=2)
    add_para(doc, "进入项目根目录，激活已安装依赖的 Python 虚拟环境。Windows 可使用 .venv\\Scripts\\Activate.ps1，Linux 或 conda 环境按本机环境管理工具执行。")
    doc.add_heading("4.3 安装或确认依赖库", level=2)
    add_para(doc, "在虚拟环境中执行 pip install -r requirements.txt。README 还要求进入 agents/reinforcement_learning 后执行 pip install -e .，以便项目内置 stable_baselines 包可被 run.py 导入。")
    doc.add_heading("4.4 配置 YAML 参数文件", level=2)
    add_para(doc, "配置文件位于 tools/cfgs/。训练前应选择与算法对应的 YAML 文件，例如 DDPG.yaml、TRPO.yaml、A2C.yaml、PPO2.yaml、SAC.yaml 或 SAC_curriculum.yaml。")
    add_picture(doc, img_paths.get("yaml"), "图4-2 YAML 配置文件界面。该图展示配置文件中的 CARLA、POLICY、GYM_ENV、RL、LOCAL_PLANNER 和 TRAFFIC_MANAGER 等字段。")
    doc.add_heading("4.5 启动软件主程序", level=2)
    add_para(doc, "训练示例命令：")
    add_para(doc, "python run.py --cfg_file=tools/cfgs/TRPO.yaml --agent_id=1 --env=CarlaGymEnv-v1")
    add_para(doc, "测试示例命令：")
    add_para(doc, "python run.py --agent_id=1 --env=CarlaGymEnv-v1 --test")
    add_para(doc, "可视化 Dashboard 示例命令：")
    add_para(doc, "python tools/pygame_dashboard.py --agent_ids 1 2 3 --window_size 100")


def add_operations(doc, img_paths):
    doc.add_heading("5 软件操作说明", level=1)
    doc.add_heading("5.1 打开软件", level=2)
    add_para(doc, "本项目的核心训练/测试入口为命令行脚本 run.py，结果查看入口为 tools/pygame_dashboard.py。若需图像化查看训练产物，启动 Dashboard 后窗口标题为 RL Visualization Dashboard，左侧显示 Active agent、Trajectory file、功能按钮和参与对比的 agent 列表。")
    add_picture(doc, img_paths.get("dashboard"), "图5-1 RL Dashboard 主界面。该界面用于选择 agent、选择轨迹文件并触发奖励曲线、轨迹图、对比图和导出操作。")
    doc.add_heading("5.2 配置参考路径参数", level=2)
    add_para(doc, "参考路径默认读取 road_maps/global_route_town04.npy。若该文件不存在，环境初始化时会在 Town04 中生成全局路线并保存到 road_maps/global_route_town04.npy。该能力由 carla_gym/envs/carla_env_v1.py 和 carla_env_v2.py 实现。")
    doc.add_heading("5.3 配置算法超参数", level=2)
    add_para(doc, "算法通过 YAML 中的 POLICY.NAME 字段选择，取值与代码确认范围一致：DDPG、TRPO、A2C、PPO2、SAC。网络类型由 POLICY.NET 配置，SAC 额外读取 SAC.GAMMA、LEARNING_RATE、BUFFER_SIZE、BATCH_SIZE 等字段。")
    doc.add_heading("5.4 配置 Frenet 规划参数", level=2)
    add_para(doc, "规划相关参数主要来自 CARLA、GYM_ENV、LOCAL_PLANNER 和 TRAFFIC_MANAGER 字段，包括 LANE_WIDTH、MAX_S、TRACK_LENGTH、TARGET_SPEED、MAX_SPEED、MIN_SPEED、MAX_SPEED、N_SPAWN_CARS 等。部分采样细节在 agents/local_planner/frenet_optimal_trajectory.py 内部实现。")
    doc.add_heading("5.5 配置车辆控制参数", level=2)
    add_para(doc, "环境初始化时创建 VehiclePIDController，当前代码中横向控制参数为 K_P=1.5、K_D=0.0、K_I=0.0；同时创建 IntelligentDriverModel 用于跟驰相关控制。")
    doc.add_heading("5.6 执行 TRPO 训练/仿真/绘图", level=2)
    add_para(doc, "将 --cfg_file 指向 tools/cfgs/TRPO.yaml 后运行 run.py。训练完成或中断后，程序在 finally 中保存最终模型；TRPO 代码还根据训练过程保存 best_* 与 step_* 模型。绘图可点击 Dashboard 中的 Plot rewards 或运行 monitor_plot.py。")
    doc.add_heading("5.7 执行 PPO2 训练/仿真/绘图", level=2)
    add_para(doc, "将 --cfg_file 指向 tools/cfgs/PPO2.yaml 后运行 run.py。PPO2 实例化时传入 model_dir，训练期间可生成 step_* 模型，最终模型保存为 PPO2_final_model.zip。")
    doc.add_heading("5.8 执行 DDPG 训练/仿真/绘图", level=2)
    add_para(doc, "将 --cfg_file 指向 tools/cfgs/DDPG.yaml 后运行 run.py。DDPG 使用 OrnsteinUhlenbeckActionNoise 和 AdaptiveParamNoiseSpec；训练过程中根据代码保存 best_*、step_* 和最终模型。")
    doc.add_heading("5.9 执行 A2C 训练/仿真/绘图", level=2)
    add_para(doc, "将 --cfg_file 指向 tools/cfgs/A2C.yaml 后运行 run.py。A2C 训练时根据 model_dir 保存 step_* 模型，最终模型保存为 A2C_final_model.zip。")
    doc.add_heading("5.10 执行 SAC 训练/仿真/绘图", level=2)
    add_para(doc, "当前代码实际支持 SAC。将 --cfg_file 指向 tools/cfgs/SAC.yaml 或 SAC_curriculum.yaml 后运行 run.py。若配置 CURRICULUM.ENABLED 为 True，或命令行使用 --use_curriculum，则程序会按 STAGES 中的 RATIO 与 N_SPAWN_CARS 分阶段调整交通密度。")
    doc.add_heading("5.11 绘制轨迹跟踪结果", level=2)
    add_para(doc, "在 Dashboard 中选择 Active agent 和 Trajectory file 后，点击 Plot trajectory 按钮生成自车轨迹或 Frenet 候选轨迹图；点击 Ego + traffic 按钮可绘制自车与周围交通参与者轨迹。")
    add_picture(doc, img_paths.get("sim2d"), "图5-2 CARLA 鸟瞰仿真视图。该图用于观察主车、车道线、周围车辆和仿真道路场景。")
    add_picture(doc, img_paths.get("trajectory"), "图5-3 轨迹跟踪结果图。该图由 Dashboard 读取 trajectory CSV 后生成，用于展示主车行驶轨迹、速度或车道变化信息。")
    doc.add_heading("5.12 绘制多算法对比分析结果", level=2)
    add_para(doc, "在 Dashboard 左侧勾选 Agents for comparison 列表中的多个 agent，点击 Plot rewards 可绘制奖励曲线，点击 Plot compare 可生成平均奖励与最高奖励对比图，点击 Artifact map 可生成训练产物概览图。")
    add_picture(doc, img_paths.get("reward"), "图5-4 多算法奖励曲线图。该图用于展示不同 agent 或算法在训练过程中的奖励变化趋势。")
    add_picture(doc, img_paths.get("compare"), "图5-5 多算法对比分析图。该图用于比较不同 agent 的尾部平均奖励、最高奖励等训练结果指标。")
    doc.add_heading("5.13 保存图片、日志、模型权重或结果文件", level=2)
    add_para(doc, "Dashboard 生成的图片默认保存到 outputs/pygame_ui/plots；预览图保存到 outputs/pygame_ui/preview。点击 Export 按钮后，程序会将选中 agent 的 monitor.csv、reproduction_info.txt、YAML 配置、models、trajectories 以及 plots 复制到 outputs/exports/export_时间戳 目录。")
    doc.add_heading("5.14 重置参数", level=2)
    add_para(doc, "当前源码未提供原说明书中“参数设置页重置按钮”的 GUI 表单。参数重置方式主要是修改 YAML 文件或重新指定命令行参数后重新运行程序；Dashboard 中的 Refresh 按钮用于刷新 agent 和轨迹文件列表。")
    doc.add_heading("5.15 退出软件", level=2)
    add_para(doc, "run.py 在训练或测试结束、异常退出或用户中断后会进入 finally 流程，保存模型并调用 env.destroy。Dashboard 中可通过窗口关闭按钮或 ESC 键退出。")


def add_io(doc):
    doc.add_heading("6 输入输出说明", level=1)
    doc.add_heading("6.1 输入参数说明", level=2)
    add_table(
        doc,
        ["参数来源", "字段或参数", "说明"],
        [
            ["命令行", "--cfg_file", "指定训练配置 YAML 文件；训练模式通常必填。"],
            ["命令行", "--env", "Gym 环境 ID，默认 CarlaGymEnv-v1。"],
            ["命令行", "--agent_id", "指定日志和模型保存目录编号，例如 logs/agent_1。"],
            ["命令行", "--num_timesteps", "训练步数，默认 1e7，run.py 会转为整数。"],
            ["命令行", "--play_mode", "显示模式：0 关闭，1 为 2D，2 为 3D；测试模式下自动开启显示。"],
            ["命令行", "--test / --test_model / --test_last", "控制模型加载测试，默认选择 best_* 或 step_* 中最新文件。"],
            ["命令行", "--carla_host / --carla_port / --tm_port / --carla_res", "CARLA 主机、端口、交通管理端口和窗口分辨率。"],
            ["命令行", "--use_curriculum / --no_curriculum", "覆盖配置文件中的课程训练开关。"],
            ["YAML", "POLICY、SAC、GYM_ENV、RL、LOCAL_PLANNER、TRAFFIC_MANAGER", "控制算法、环境、奖励、规划速度和交通密度。"],
        ],
        widths=[1.0, 2.1, 3.3],
    )
    doc.add_heading("6.2 配置文件说明", level=2)
    add_para(doc, "配置文件目录为 tools/cfgs/，当前确认存在 A2C.yaml、DDPG.yaml、PPO2.yaml、TRPO.yaml、SAC.yaml、SAC_curriculum.yaml、SAC_explore.yaml、SAC_safe.yaml。")
    doc.add_heading("6.3 运行日志说明", level=2)
    add_para(doc, "训练模式下，Monitor 将 episode 奖励、长度和时间写入 logs/agent_x/monitor.csv。run.py 还写入 reproduction_info.txt，包含 Git commit id、程序参数和配置内容。")
    doc.add_heading("6.4 模型权重文件说明", level=2)
    add_para(doc, "模型保存在 logs/agent_x/models/。最终模型命名为 算法名_final_model.zip；训练中间模型可按算法保存为 best_*.zip 或 step_*.zip。测试时如未指定 --test_model，程序会在模型目录中选择 best_* 或 step_* 的最新编号。")
    doc.add_heading("6.5 图像输出说明", level=2)
    add_para(doc, "Dashboard 生成的图像使用 matplotlib 保存为 PNG，文件名包括 rewards_*.png、trajectory_*.png、surrounding_tracks_*.png、comparison_*.png、artifact_overview_*.png。monitor_plot.py 默认输出 monitor_plot.png，也可通过 --output_path 指定路径。未在源码中确认支持 jpg、bmp、tiff 等格式，因此本文档按 PNG 输出说明。")
    doc.add_heading("6.6 结果数据文件说明", level=2)
    add_table(
        doc,
        ["文件或目录", "生成位置", "内容"],
        [
            ["monitor.csv", "logs/agent_x/", "训练 episode 奖励 r、长度 l、时间 t 等监控数据。"],
            ["reproduction_info.txt", "logs/agent_x/", "Git 提交、启动参数和配置快照。"],
            ["*.yaml 配置副本", "logs/agent_x/", "本次训练使用的配置文件副本。"],
            ["trajectory_ep*_*.csv", "logs/agent_x/trajectories/", "自车位置、Frenet s/d、速度、变道标记、前车信息。"],
            ["plans_ep*_*.csv / latest_plans.csv", "logs/agent_x/trajectories/plans/", "候选轨迹和选定轨迹的 x/y 点。"],
            ["surrounding_ep*_*.csv / latest_surrounding_tracks.csv", "logs/agent_x/trajectories/surrounding/", "自车与周围交通车辆的轨迹采样。"],
            ["*.zip 模型", "logs/agent_x/models/", "Stable Baselines 模型权重和参数。"],
            ["*.png 图像", "outputs/pygame_ui/plots/", "Dashboard 生成的奖励、轨迹、对比和产物概览图。"],
        ],
        widths=[1.9, 1.9, 2.6],
    )


def add_faq(doc):
    doc.add_heading("7 异常处理与常见问题", level=1)
    rows = [
        ["CARLA 服务端连接失败", "run.py 输出连接或超时相关错误。", "CARLA 未启动、主机或端口不一致。", "先启动 CARLA；确认 --carla_host、--carla_port 与服务端 world-port 一致。"],
        ["端口被占用", "CARLA 或客户端无法绑定 2000/8000 端口。", "已有 CARLA 或 Traffic Manager 进程占用端口。", "关闭旧进程，或修改 --carla_port、--tm_port 后重新启动。"],
        ["Python 依赖缺失", "导入 gym、tensorflow、pygame、yaml 等失败。", "未安装 requirements.txt 或未激活正确虚拟环境。", "激活虚拟环境，执行 pip install -r requirements.txt，并安装 agents/reinforcement_learning。"],
        ["CUDA/GPU 不可用", "CARLA 帧率低或 TensorFlow 使用 CPU。", "显卡驱动、CUDA 或 GPU 资源不可用。", "降低 CARLA 分辨率和画质；检查本机 GPU 驱动与 TensorFlow 1.14 兼容性。"],
        ["YAML 配置文件格式错误", "cfg_from_yaml_file 抛出解析异常。", "缩进、冒号或字段类型错误。", "使用 tools/cfgs 中现有文件为模板修改，避免 Tab 和非法层级。"],
        ["配置路径不存在", "run.py 无法打开 --cfg_file。", "路径拼写错误或工作目录不在项目根目录。", "切换到项目根目录，使用 tools/cfgs/xxx.yaml 相对路径。"],
        ["模型文件不存在", "测试模式无法找到 best_* 或 step_*。", "agent_id 目录没有训练生成的模型。", "检查 logs/agent_x/models；指定 --test_model 或先完成训练。"],
        ["日志目录已存在或无写入权限", "训练开始时 os.mkdir 报错。", "同一 agent_id 目录已存在，或目录权限不足。", "更换 --agent_id，或确认目录可写；必要时人工清理旧实验目录。"],
        ["图片导出失败", "Dashboard 状态栏提示保存失败或无图像。", "outputs 目录不可写、缺少 CSV 或缺少 matplotlib。", "确认输出目录权限，先运行训练/测试生成 monitor.csv 和 trajectory CSV。"],
        ["训练过程中断", "终端中断或异常退出。", "用户中断、CARLA 崩溃、配置不匹配。", "run.py 会尝试在 finally 中保存最终模型并销毁环境；重启 CARLA 后可重新训练。"],
        ["软件无法正常退出", "窗口关闭后 CARLA 仍有残留 Actor。", "环境销毁流程未完成或外部强制结束。", "优先使用正常退出；若服务端仍占用资源，重启 CARLA 服务端。"],
    ]
    add_table(doc, ["问题", "现象", "可能原因", "处理方法"], rows, widths=[1.25, 1.55, 1.55, 2.05])


def add_appendix(doc, img_paths):
    doc.add_heading("8 附录", level=1)
    doc.add_heading("8.1 项目目录结构", level=2)
    add_table(
        doc,
        ["目录/文件", "说明"],
        [
            [".venv/", "本地 Python 虚拟环境目录。"],
            ["agents/local_planner/", "Frenet 轨迹规划和样条路径模块。"],
            ["agents/low_level_controller/", "车辆 PID 控制与 IDM 跟驰模型。"],
            ["agents/reinforcement_learning/", "项目内置 Stable Baselines 强化学习库。"],
            ["carla_gym/envs/", "CARLA Gym 环境、场景文件和环境注册入口。"],
            ["logs/", "训练日志、模型、轨迹和配置副本。"],
            ["outputs/", "Dashboard 图像、预览和导出归档。"],
            ["road_maps/", "全局路线 npy 与道路地图 png。"],
            ["tools/cfgs/", "算法和环境 YAML 配置文件。"],
            ["tools/modules.py", "CARLA 世界、交通流、HUD、输入和渲染模块封装。"],
            ["tools/pygame_dashboard.py", "Pygame/Tk 可视化 Dashboard。"],
            ["monitor_plot.py", "奖励曲线命令行绘图脚本。"],
            ["run.py", "软件训练和测试主入口。"],
        ],
        widths=[2.2, 4.2],
    )
    doc.add_heading("8.2 命令行参数说明", level=2)
    add_table(
        doc,
        ["参数", "默认值", "说明"],
        [
            ["--cfg_file", "None", "训练配置文件路径；测试时若未提供则从 logs/agent_x 中读取 YAML。"],
            ["--env", "CarlaGymEnv-v1", "Gym 环境 ID。"],
            ["--log_interval", "100", "模型训练日志间隔。"],
            ["--agent_id", "None", "日志和模型保存编号。"],
            ["--num_timesteps", "1e7", "训练步数。"],
            ["--play_mode", "0", "渲染显示模式：0 关闭，1 2D，2 3D。"],
            ["--verbosity", "0", "终端输出详细程度。"],
            ["--test", "False", "启用模型加载测试模式。"],
            ["--carla_host", "127.0.0.1", "CARLA 服务端主机地址。"],
            ["--carla_port", "2000", "CARLA 服务端端口。"],
            ["--tm_port", "8000", "Traffic Manager 端口。"],
            ["--backend", "pygame", "Dashboard 参数；可选 pygame 或 tk。"],
        ],
        widths=[1.7, 1.2, 3.5],
    )
    doc.add_heading("8.3 YAML 配置字段说明", level=2)
    add_table(
        doc,
        ["字段组", "主要字段", "说明"],
        [
            ["CARLA", "DT、LANE_WIDTH、MAX_S", "仿真时间步、车道宽度和全局 Frenet 距离上限。"],
            ["POLICY", "NAME、NET、CNN_EXTRACTOR、ACTION_NOISE、PARAM_NOISE_STD", "算法名称、网络类型、CNN 特征提取器和 DDPG 噪声参数。"],
            ["SAC", "GAMMA、LEARNING_RATE、BUFFER_SIZE、BATCH_SIZE 等", "SAC 算法专用超参数；仅 SAC 配置文件包含。"],
            ["CURRICULUM", "ENABLED、STAGES、RATIO、N_SPAWN_CARS", "课程训练阶段配置；可由命令行开关覆盖。"],
            ["GYM_ENV", "TRACK_LENGTH、TARGET_SPEED、LOOK_BACK、TIME_STEP 等", "Gym 环境状态表示、目标速度和回看步数等参数。"],
            ["RL", "W_SPEED、W_R_SPEED、LANE_CHANGE_REWARD、COLLISION 等", "奖励与惩罚权重。"],
            ["LOCAL_PLANNER", "MIN_SPEED、MAX_SPEED", "局部规划速度范围。"],
            ["TRAFFIC_MANAGER", "N_SPAWN_CARS、MIN_SPEED、MAX_SPEED", "随机交通车辆数量和速度范围。"],
        ],
        widths=[1.3, 2.4, 2.7],
    )
    doc.add_heading("8.4 界面截图说明", level=2)
    add_table(
        doc,
        ["图号", "图名", "对应功能"],
        [
            ["图1-1", "软件运行流程示意图", "说明端到端运行步骤和输出位置。"],
            ["图3-1", "软件总体架构示意图", "说明源码模块与功能模块对应关系。"],
            ["图4-1", "CARLA 仿真服务端界面", "说明 CARLA 服务端运行状态。"],
            ["图4-2", "YAML 配置文件界面", "说明配置字段修改位置。"],
            ["图5-1", "RL Dashboard 主界面", "说明可视化操作入口和按钮。"],
            ["图5-2", "CARLA 鸟瞰仿真视图", "说明 2D 渲染和交通场景。"],
            ["图5-3", "轨迹跟踪结果图", "说明轨迹 CSV 绘图结果。"],
            ["图5-4", "多算法奖励曲线图", "说明 monitor.csv 奖励曲线。"],
            ["图5-5", "多算法对比分析图", "说明多 agent 指标对比。"],
        ],
        widths=[1.0, 2.2, 3.2],
    )
    doc.add_heading("8.5 版本修订记录", level=2)
    add_table(
        doc,
        ["版本号", "修订日期", "修订人", "修订内容"],
        [
            ["V1.0", "2026-04-26", "彭靖", "初始版本，包含仿真环境连接、参数配置、训练执行、轨迹绘制、结果保存等功能说明。"],
            ["V1.0 修订版", "2026-05-26", "彭靖/文档修订辅助", "根据当前项目代码修订软件名称、运行方式、算法范围、功能模块、输入输出、异常处理和图题说明。"],
        ],
        widths=[1.2, 1.2, 1.35, 2.65],
    )
    doc.add_heading("8.6 修订摘要与需人工确认事项", level=2)
    add_table(
        doc,
        ["项目", "修订结果"],
        [
            ["统一后的软件名称", SOFTWARE_NAME],
            ["删除或压缩的内容", "压缩自动驾驶产业背景、论文式算法意义、宣传式和绝对化描述；删除原稿中未在代码确认的表单页、保存对话框和 jpg/bmp/tiff 输出描述。"],
            ["新增章节", "运行环境、软件结构与功能模块、安装与启动、输入输出说明、异常处理、附录目录结构、命令行参数和 YAML 字段说明。"],
            ["新增功能模块说明", "SAC 训练、课程训练配置、Pygame/Tk Dashboard、周围交通轨迹绘制、Artifact map、Export 导出归档。"],
            ["代码确认的算法列表", "DDPG、TRPO、A2C、PPO2、SAC。"],
            ["代码确认的主入口文件", "run.py；绘图入口为 monitor_plot.py 和 tools/pygame_dashboard.py。"],
            ["代码确认的配置文件目录", "tools/cfgs/。"],
            ["代码确认的输出目录", "logs/agent_x/、outputs/pygame_ui/、outputs/exports/。"],
            ["仍需人工确认事项", "本机 CARLA 0.9.9.2 安装路径、实际 GPU/CUDA 版本、最终提交软著时的软件著作权归属信息、是否需要将修订版日期调整为正式提交日期。"],
        ],
        widths=[1.7, 4.7],
    )


def build():
    ensure_dirs()
    extract_original_images()
    arch_path, flow_path = make_diagrams()

    img_paths = {
        "carla": ORIG_IMG_DIR / "image2.png",
        "sim2d": ORIG_IMG_DIR / "image6.png",
        "yaml": ORIG_IMG_DIR / "image4.png",
        "dashboard": ORIG_IMG_DIR / "image5.png",
        "reward": ORIG_IMG_DIR / "image7.png",
        "compare": ORIG_IMG_DIR / "image8.png",
    }
    trajectory_candidates = sorted((ROOT / "outputs" / "pygame_ui" / "plots").glob("trajectory_*.png"))
    if trajectory_candidates:
        img_paths["trajectory"] = trajectory_candidates[-1]
    else:
        img_paths["trajectory"] = ORIG_IMG_DIR / "image3.png"

    doc = Document()
    setup_document(doc)
    add_cover(doc)
    doc.add_heading("目录", level=1)
    add_para(doc, "本文档按照软件概述、运行环境、软件结构与功能模块、安装与启动、操作说明、输入输出、异常处理和附录组织。Word 打开后可根据标题样式自动生成或更新目录。")
    doc.add_page_break()
    add_overview(doc, flow_path)
    add_environment(doc)
    add_structure(doc, arch_path)
    add_installation(doc, img_paths)
    add_operations(doc, img_paths)
    add_io(doc)
    add_faq(doc)
    add_appendix(doc, img_paths)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
